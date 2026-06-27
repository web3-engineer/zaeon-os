"""
ABNT NBR 14724:2011 — Motor de geração de documentos Word
Usa python-docx para controle preciso de toda a formatação.
"""
import base64
import io
import re
import urllib.request
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER,
)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.part import Part
from docx.opc.packuri import PackURI

import abnt_rules as R
from models import DocumentPayload


def build_abnt_document(payload: DocumentPayload) -> bytes:
    """
    Constrói o documento Word completo conforme ABNT NBR 14724:2011.
    Retorna os bytes do .docx.
    """
    doc = Document()
    font = payload.font_choice or R.FONT_BODY

    _setup_page(doc)
    _setup_base_style(doc, font)

    # ── Seção 1: CAPA (não contada, sem numeração) ──
    _add_cover_page(doc, payload, font)

    # ── Seção 2: PRÉ-TEXTUAIS (contadas a partir da folha de rosto = pág. 1) ──
    pretextual = _new_section(doc)
    _set_section_start_page(pretextual, 1)  # capa não conta; folha de rosto = 1

    _add_title_page(doc, payload, font)
    doc.add_page_break()

    _add_resumo(doc, payload.resumo, payload.palavras_chave, font)
    doc.add_page_break()

    if payload.abstract and payload.abstract.strip():
        _add_abstract(doc, payload.abstract, payload.keywords, font)
        doc.add_page_break()

    # Listas de ilustrações/tabelas (quando houver) — antes do sumário
    illustrations, tables = _collect_floats(payload.sections)
    if illustrations:
        _add_list_of_illustrations(doc, illustrations, font)
        doc.add_page_break()
    if tables:
        _add_list_of_tables(doc, tables, font)
        doc.add_page_break()

    _add_table_of_contents(doc, payload.sections, bool(payload.references),
                           payload.apendices, payload.anexos, font)

    # ── Seção 3: TEXTUAL + PÓS-TEXTUAL (numeração visível, contagem contínua) ──
    textual = _new_section(doc)
    _set_section_continue(textual)  # não reinicia (continua da folha de rosto)

    counters = {"figura": 0, "year": payload.cover.year if payload.cover else None}
    for i, section in enumerate(payload.sections, start=1):
        _add_body_section(doc, section, i, font, counters)

    doc.add_page_break()
    _add_references(doc, payload.references, font, counters)

    _add_appendices(doc, payload.apendices, "APÊNDICE", R.BM_APPENDIX, font, counters)
    _add_appendices(doc, payload.anexos, "ANEXO", R.BM_ANNEX, font, counters)

    _add_page_numbers(doc)

    if counters.get("footnotes"):
        _build_footnotes_part(doc, counters["footnotes"], font)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── Configuração de página ─────────────────────────────────────

def _setup_section(section):
    section.page_width  = Cm(R.PAGE_WIDTH_CM)
    section.page_height = Cm(R.PAGE_HEIGHT_CM)
    section.left_margin   = Cm(R.MARGIN_LEFT_CM)
    section.right_margin  = Cm(R.MARGIN_RIGHT_CM)
    section.top_margin    = Cm(R.MARGIN_TOP_CM)
    section.bottom_margin = Cm(R.MARGIN_BOTTOM_CM)


def _new_section(doc: Document):
    """Quebra de seção (nova página) com a mesma configuração de página ABNT."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    _setup_section(section)
    return section


def _set_section_start_page(section, start: int):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), str(start))


def _set_section_continue(section):
    """Remove o 'start' herdado (add_section copia o sectPr anterior) para a
    numeração continuar em vez de reiniciar."""
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is not None:
        if qn('w:start') in pgNumType.attrib:
            del pgNumType.attrib[qn('w:start')]
        if not pgNumType.attrib:
            sectPr.remove(pgNumType)


def _setup_page(doc: Document):
    section = doc.sections[0]
    section.page_width  = Cm(R.PAGE_WIDTH_CM)
    section.page_height = Cm(R.PAGE_HEIGHT_CM)
    section.left_margin   = Cm(R.MARGIN_LEFT_CM)
    section.right_margin  = Cm(R.MARGIN_RIGHT_CM)
    section.top_margin    = Cm(R.MARGIN_TOP_CM)
    section.bottom_margin = Cm(R.MARGIN_BOTTOM_CM)


def _setup_base_style(doc: Document, font_name: str):
    style = doc.styles['Normal']
    style.font.name = font_name
    style.font.size = Pt(R.FONT_SIZE_BODY_PT)
    fmt = style.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = R.LINE_SPACING_BODY
    fmt.first_line_indent = Cm(R.PARAGRAPH_INDENT_CM)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # corpo justificado (NBR 14724)
    fmt.space_before = Pt(0)
    fmt.space_after  = Pt(0)


# ── Capa ───────────────────────────────────────────────────────

def _add_cover_page(doc: Document, payload: DocumentPayload, font: str):
    cover = payload.cover

    def centered(text: str, size: int = R.FONT_SIZE_BODY_PT, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = bold

    # Instituição — topo centralizado
    centered(cover.institution.upper(), bold=True)
    centered(cover.course)

    # Espaço entre instituição e autor
    for _ in range(8):
        _blank_line(doc, font)

    centered(cover.author.upper())

    for _ in range(8):
        _blank_line(doc, font)

    # Título em maiúsculas, negrito
    centered(payload.title.upper(), bold=True)

    # Natureza do trabalho (se TCC ou Relatório)
    if payload.work_type in ("TCC", "Relatório"):
        for _ in range(4):
            _blank_line(doc, font)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(f"{payload.work_type} apresentado como requisito parcial.")
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_BODY_PT)
        if cover.advisor:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.first_line_indent = Cm(0)
            run2 = p2.add_run(f"Orientador: {cover.advisor}")
            run2.font.name = font
            run2.font.size = Pt(R.FONT_SIZE_BODY_PT)

    # Cidade e ano — rodapé
    for _ in range(8):
        _blank_line(doc, font)

    centered(cover.city.upper())
    centered(cover.year)


# ── Folha de rosto ─────────────────────────────────────────────

def _add_title_page(doc: Document, payload: DocumentPayload, font: str):
    cover = payload.cover

    def centered(text: str, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_BODY_PT)
        run.bold = bold

    centered(cover.author.upper())

    for _ in range(8):
        _blank_line(doc, font)

    centered(payload.title.upper(), bold=True)

    for _ in range(6):
        _blank_line(doc, font)

    # Nota explicativa (alinhada à direita, metade da largura)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(8)
    run = p.add_run(
        f"{payload.work_type} apresentado ao Curso de {cover.course} "
        f"da {cover.institution}, como requisito parcial para conclusão."
    )
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_BODY_PT)

    if cover.advisor:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p2.paragraph_format.first_line_indent = Cm(0)
        p2.paragraph_format.left_indent = Cm(8)
        run2 = p2.add_run(f"Orientador(a): {cover.advisor}")
        run2.font.name = font
        run2.font.size = Pt(R.FONT_SIZE_BODY_PT)

    for _ in range(8):
        _blank_line(doc, font)

    centered(cover.city.upper())
    centered(cover.year)


# ── Resumo ─────────────────────────────────────────────────────

def _add_resumo(doc: Document, resumo: str, palavras_chave: list, font: str):
    _section_title(doc, "RESUMO", font)

    p = doc.add_paragraph(resumo)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # resumo justificado (NBR 14724)
    p.paragraph_format.first_line_indent = Cm(0)  # resumo sem recuo
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = R.LINE_SPACING_BODY
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_BODY_PT)

    _blank_line(doc, font)

    _keywords_line(doc, "Palavras-chave: ", palavras_chave, font)


# ── Abstract (resumo em língua estrangeira) ────────────────────

def _add_abstract(doc: Document, abstract: str, keywords: list, font: str):
    _section_title(doc, "ABSTRACT", font)

    p = doc.add_paragraph(abstract)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = R.LINE_SPACING_BODY
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_BODY_PT)

    _blank_line(doc, font)

    _keywords_line(doc, "Keywords: ", keywords or [], font)


def _keywords_line(doc: Document, label: str, terms: list, font: str):
    kw = doc.add_paragraph()
    kw.alignment = WD_ALIGN_PARAGRAPH.LEFT  # palavras-chave alinhadas à esquerda
    kw.paragraph_format.first_line_indent = Cm(0)
    r_bold = kw.add_run(label)
    r_bold.bold = True
    r_bold.font.name = font
    r_bold.font.size = Pt(R.FONT_SIZE_BODY_PT)
    r_kw = kw.add_run("; ".join(terms) + ("." if terms else ""))
    r_kw.font.name = font
    r_kw.font.size = Pt(R.FONT_SIZE_BODY_PT)


# ── Sumário (NBR 6027) — entradas com PAGEREF (números reais no PDF) ──

def _add_table_of_contents(doc: Document, sections: list, has_refs: bool,
                           apendices: list, anexos: list, font: str):
    """Sumário com entradas (nº + título), líder pontilhado e número de página
    via campo PAGEREF a um bookmark em cada título. LibreOffice/Word resolvem o
    PAGEREF no layout — o PDF headless já mostra os números de página corretos."""
    _section_title(doc, "SUMÁRIO", font)

    entries = []  # (nível, texto, bookmark)
    for i, sec in enumerate(sections, start=1):
        entries.append((0, f"{i} {sec.title.upper()}", f"{R.BM_SECTION}{i}"))
        for j, sub in enumerate(getattr(sec, "subsections", []) or [], start=1):
            entries.append((1, f"{i}.{j} {sub.title.upper()}", f"{R.BM_SECTION}{i}sub{j}"))
    if has_refs:
        entries.append((0, R.REFERENCES_TITLE, R.BM_REFERENCES))
    for idx, ap in enumerate(apendices or []):
        entries.append((0, f"APÊNDICE {chr(65 + idx)}{R.FIGURE_CAPTION_DELIM}{ap.title.upper()}",
                        f"{R.BM_APPENDIX}{idx}"))
    for idx, ax in enumerate(anexos or []):
        entries.append((0, f"ANEXO {chr(65 + idx)}{R.FIGURE_CAPTION_DELIM}{ax.title.upper()}",
                        f"{R.BM_ANNEX}{idx}"))

    for level, text, bm in entries:
        _index_entry(doc, text, bm, font, level)
    _enable_update_fields(doc)


def _add_appendices(doc: Document, items: list, label: str, bm_prefix: str,
                    font: str, counters: dict):
    """Emite apêndices/anexos pós-textuais: 'APÊNDICE A – TÍTULO' (nova página) +
    conteúdo. Cada um vira alvo de bookmark para aparecer no sumário."""
    for idx, item in enumerate(items or []):
        doc.add_page_break()
        title = f"{label} {chr(65 + idx)}{R.FIGURE_CAPTION_DELIM}{item.title.upper()}"
        tp = _section_title(doc, title, font, outline_level=0)
        if counters is not None:
            _add_bookmark(tp, f"{bm_prefix}{idx}", counters)
        for block in _split_paragraphs(item.content):
            _add_paragraph(doc, block, font, counters)


# ── Listas de Ilustrações e Tabelas (NBR 14724) ────────────────

def _collect_floats(sections: list):
    """Pré-varre as seções (na ordem do texto) coletando figuras/quadros (→ Lista
    de Ilustrações) e tabelas (→ Lista de Tabelas). A numeração espelha exatamente
    a da renderização do corpo (mesmos contadores por tipo)."""
    illustrations, tables = [], []
    c = {"figura": 0, "tabela": 0, "quadro": 0}

    def scan(content: str):
        for block in _split_paragraphs(content or ""):
            if _is_figure(block):
                m = _FIGURE_RE.match(block.strip())
                title = m.group("alt").partition("|")[0].strip()
                c["figura"] += 1
                illustrations.append((R.FIGURE_LABEL, c["figura"], title, f"{R.BM_FIGURE}{c['figura']}"))
            elif _is_table_block(block):
                kind, title, _fonte, rows = _parse_table_block(block)
                if not rows:
                    continue
                if kind == "tabela":
                    c["tabela"] += 1
                    tables.append((R.TABLE_LABEL, c["tabela"], title, f"{R.BM_TABLE}{c['tabela']}"))
                else:
                    c["quadro"] += 1
                    illustrations.append((R.QUADRO_LABEL, c["quadro"], title, f"{R.BM_QUADRO}{c['quadro']}"))

    for sec in sections:
        scan(sec.content)
        for sub in getattr(sec, "subsections", []) or []:
            scan(sub.content)
    return illustrations, tables


def _add_list_of_illustrations(doc: Document, items: list, font: str):
    _section_title(doc, "LISTA DE ILUSTRAÇÕES", font)
    for label, num, title, bm in items:
        _index_entry(doc, f"{label} {num}{R.FIGURE_CAPTION_DELIM}{title}", bm, font)
    _enable_update_fields(doc)


def _add_list_of_tables(doc: Document, items: list, font: str):
    _section_title(doc, "LISTA DE TABELAS", font)
    for label, num, title, bm in items:
        _index_entry(doc, f"{label} {num}{R.FIGURE_CAPTION_DELIM}{title}", bm, font)
    _enable_update_fields(doc)


def _index_entry(doc: Document, text: str, bookmark: str, font: str, level: int = 0):
    """Linha de índice: 'texto ......... <página>' (tab à direita + líder pontilhado)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if level > 0:
        p.paragraph_format.left_indent = Cm(0.5 * level)
    usable = R.PAGE_WIDTH_CM - R.MARGIN_LEFT_CM - R.MARGIN_RIGHT_CM  # 16 cm (margem direita)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(usable), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r = p.add_run(text + "\t")
    r.font.name = font
    r.font.size = Pt(R.FONT_SIZE_BODY_PT)
    _add_field(p, f' PAGEREF {bookmark} \\h ', "0", font)


def _add_field(paragraph, instr: str, placeholder: str, font: str, bold: bool = False):
    """Campo Word: fldChar(begin) instrText fldChar(separate) <placeholder> fldChar(end)."""
    def fld_run():
        run = paragraph.add_run()
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_BODY_PT)
        run.bold = bold
        return run._r
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); fld_run().append(begin)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; fld_run().append(it)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); fld_run().append(sep)
    ph = paragraph.add_run(placeholder)
    ph.font.name = font
    ph.font.size = Pt(R.FONT_SIZE_BODY_PT)
    ph.bold = bold
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); fld_run().append(end)


def _add_bookmark(paragraph, name: str, counters: dict):
    """Marca o parágrafo com um bookmark (alvo de PAGEREF no sumário/listas)."""
    bid = counters.get("bmid", 0) + 1
    counters["bmid"] = bid
    p_elem = paragraph._p
    pPr = p_elem.find(qn('w:pPr'))
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bid))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bid))
    if pPr is not None:
        pPr.addnext(start)   # bookmarkStart logo após o pPr (ordem válida)
    else:
        p_elem.insert(0, start)
    p_elem.append(end)


def _enable_update_fields(doc: Document):
    """Marca os campos do documento para atualização ao abrir (popula o TOC com
    os números de página). Inserido na posição válida do CT_Settings."""
    settings = doc.settings.element
    el = settings.find(qn('w:updateFields'))
    if el is None:
        el = OxmlElement('w:updateFields')
        # w:updateFields precede compat/rsids/etc. no schema CT_Settings
        anchor = None
        for tag in ('w:compat', 'w:rsids', 'w:mathPr', 'w:themeFontLang'):
            found = settings.find(qn(tag))
            if found is not None:
                anchor = found
                break
        if anchor is not None:
            anchor.addprevious(el)
        else:
            settings.append(el)
    el.set(qn('w:val'), 'true')


def _set_outline_level(paragraph, level: int):
    """Define o nível de tópico do parágrafo (entra no sumário sem alterar o
    visual ABNT do título)."""
    pPr = paragraph._p.get_or_add_pPr()
    el = pPr.find(qn('w:outlineLvl'))
    if el is None:
        el = OxmlElement('w:outlineLvl')
        # outlineLvl precede rPr/sectPr/pPrChange no schema CT_PPr
        anchor = None
        for tag in ('w:rPr', 'w:sectPr', 'w:pPrChange'):
            found = pPr.find(qn(tag))
            if found is not None:
                anchor = found
                break
        if anchor is not None:
            anchor.addprevious(el)
        else:
            pPr.append(el)
    el.set(qn('w:val'), str(level))


# ── Seções do corpo ────────────────────────────────────────────

def _add_body_section(doc: Document, section_data, number: int, font: str, counters: dict = None):
    title_text = f"{number} {section_data.title.upper()}"
    tp = _section_title(doc, title_text, font, outline_level=0)
    if counters is not None:
        _add_bookmark(tp, f"{R.BM_SECTION}{number}", counters)

    for block in _split_paragraphs(section_data.content):
        _add_paragraph(doc, block, font, counters)

    for j, sub in enumerate(section_data.subsections, start=1):
        sub_title = f"{number}.{j} {sub.title.upper()}"
        stp = _subsection_title(doc, sub_title, font)
        if counters is not None:
            _add_bookmark(stp, f"{R.BM_SECTION}{number}sub{j}", counters)
        for block in _split_paragraphs(sub.content):
            _add_paragraph(doc, block, font, counters)


def _add_paragraph(doc: Document, text: str, font: str, counters: dict = None):
    if not text.strip():
        return
    if counters is not None and _is_table_block(text):
        _add_table_block(doc, text, font, counters)
        return
    if counters is not None and _is_figure(text):
        _add_figure(doc, text, font, counters)
        return
    if _is_alinea_block(text):
        _add_alineas(doc, text, font)
        return
    if _is_long_quote(text):
        _add_long_quote(doc, text, font)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY  # corpo justificado (NBR 14724)
        p.paragraph_format.first_line_indent = Cm(R.PARAGRAPH_INDENT_CM)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = R.LINE_SPACING_BODY
        _add_text_with_footnotes(p, text, font, counters)


def _add_long_quote(doc: Document, text: str, font: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(R.LONG_QUOTE_INDENT_CM)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(R.SPACE_BEFORE_TITLE_PT)
    p.paragraph_format.space_after  = Pt(R.SPACE_AFTER_TITLE_PT)
    run = p.add_run(_strip_quotes(text))  # citação longa não leva aspas (NBR 10520)
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_QUOTE_PT)


_QUOTE_CHARS_OPEN  = ('"', '“', '‘', '«')
_QUOTE_CHARS_CLOSE = ('"', '”', '’', '»')


def _is_quotation(text: str) -> bool:
    return text.strip()[:1] in _QUOTE_CHARS_OPEN


def _strip_quotes(text: str) -> str:
    t = text.strip()
    if t[:1] in _QUOTE_CHARS_OPEN:
        t = t[1:]
    if t[-1:] in _QUOTE_CHARS_CLOSE:
        t = t[:-1]
    return t.strip()


def _estimate_lines(text: str) -> int:
    """Estima o número de linhas renderizadas (largura útil ≈ 90 chars/linha)."""
    import math
    total = 0
    for line in text.split("\n"):
        total += max(1, math.ceil(len(line) / R.CHARS_PER_LINE))
    return total


def _is_long_quote(text: str) -> bool:
    """Citação direta longa = é uma citação (marcada por aspas) E ultrapassa 3
    linhas (NBR 10520). O marcador de aspas evita que prosa comum longa seja
    tratada como bloco de citação."""
    return _is_quotation(text) and _estimate_lines(text.strip()) > R.LONG_QUOTE_MIN_LINES


def _split_paragraphs(content: str) -> list:
    return [b.strip() for b in content.split("\n\n") if b.strip()]


# ── Alíneas (NBR 6024) — lista markdown "- item" → a) b) c) ────

_ALINEA_LINE_RE = re.compile(r'^[-*]\s+\S')


def _is_alinea_block(text: str) -> bool:
    lines = [l.strip() for l in text.strip().split("\n") if l.strip()]
    return len(lines) >= 2 and all(_ALINEA_LINE_RE.match(l) for l in lines)


def _add_alineas(doc: Document, text: str, font: str):
    """Renderiza alíneas NBR 6024: a) b) c) minúsculas, recuo, ';' em cada e '.'
    na última."""
    items = [re.sub(r'^[-*]\s+', '', l.strip())
             for l in text.strip().split("\n") if l.strip()]
    for idx, item in enumerate(items):
        letter = chr(ord('a') + idx) if idx < 26 else f"a{idx}"
        content = item.rstrip('.;, ')
        punct = '.' if idx == len(items) - 1 else ';'
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(R.PARAGRAPH_INDENT_CM)  # recuo das alíneas
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = R.LINE_SPACING_BODY
        r = p.add_run(f"{letter}) {content}{punct}")
        r.font.name = font
        r.font.size = Pt(R.FONT_SIZE_BODY_PT)


# ── Notas de rodapé (NBR 14724) — marcador inline ^[nota] ──────

_FOOTNOTE_RE = re.compile(r'\^\[(.+?)\]', re.DOTALL)


def _add_text_with_footnotes(p, text: str, font: str, counters: dict = None):
    """Adiciona o texto do parágrafo, convertendo marcadores ^[nota] em notas
    de rodapé reais (referência sobrescrita no corpo + nota na parte footnotes)."""
    def body_run(s):
        r = p.add_run(s)
        r.font.name = font
        r.font.size = Pt(R.FONT_SIZE_BODY_PT)

    if counters is None or '^[' not in text:
        body_run(text)
        return

    pos = 0
    for m in _FOOTNOTE_RE.finditer(text):
        if text[pos:m.start()]:
            body_run(text[pos:m.start()])
        _add_footnote_ref(p, m.group(1).strip(), font, counters)
        pos = m.end()
    if text[pos:]:
        body_run(text[pos:])


def _add_footnote_ref(p, note: str, font: str, counters: dict):
    notes = counters.setdefault("footnotes", [])
    notes.append(note)
    fid = len(notes)
    run = p.add_run()
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    va = OxmlElement('w:vertAlign')
    va.set(qn('w:val'), 'superscript')
    rpr.append(va)
    ref = OxmlElement('w:footnoteReference')
    ref.set(qn('w:id'), str(fid))
    run._r.append(ref)


def _xml_escape(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _build_footnotes_part(doc: Document, notes: list, font: str):
    """Cria a parte word/footnotes.xml (10pt, espaço simples) e a relaciona ao
    documento. Inclui os separators reservados (id -1 e 0); notas começam em 1."""
    sz = R.FONT_SIZE_NOTE_PT * 2  # meios-pontos
    items = [
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>',
    ]
    for i, note in enumerate(notes, start=1):
        items.append(
            f'<w:footnote w:id="{i}"><w:p>'
            f'<w:pPr><w:spacing w:line="240" w:lineRule="auto"/></w:pPr>'
            f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:rPr><w:rFonts w:ascii="{_xml_escape(font)}" w:hAnsi="{_xml_escape(font)}"/>'
            f'<w:sz w:val="{sz}"/></w:rPr>'
            f'<w:t xml:space="preserve"> {_xml_escape(note)}</w:t></w:r>'
            f'</w:p></w:footnote>'
        )
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        + "".join(items) +
        '</w:footnotes>'
    )
    part = Part(PackURI('/word/footnotes.xml'), R.FOOTNOTES_CONTENT_TYPE,
                xml.encode('utf-8'), doc.part.package)
    doc.part.relate_to(part, R.FOOTNOTES_REL_TYPE)


# ── Figuras / Ilustrações (NBR 14724:2024) ─────────────────────
# Marcador (markdown estendido):  ![Título|Fonte](url)
#   - "Fonte" é opcional; "própria"/"autor" → "elaborado pelo próprio autor".
_FIGURE_RE = re.compile(r'^!\[(?P<alt>.*?)\]\((?P<url>\S+?)\)$', re.DOTALL)
_OWN_WORK_TOKENS = {
    "própria", "propria", "autor", "autoral", "own",
    "elaboração própria", "elaboracao propria", "elaborado pelo autor",
    "elaborado pelo próprio autor", "elaborado pelo proprio autor",
}


def _is_figure(text: str) -> bool:
    return bool(_FIGURE_RE.match(text.strip()))


def _add_figure(doc: Document, text: str, font: str, counters: dict):
    m = _FIGURE_RE.match(text.strip())
    alt, url = m.group("alt"), m.group("url")
    title, _, fonte_raw = alt.partition("|")
    counters["figura"] = counters.get("figura", 0) + 1
    number = counters["figura"]

    # Legenda no TOPO: "Figura N – Título" (travessão, sem ponto final), 10pt, centralizada
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cap.paragraph_format.space_before = Pt(R.SPACE_BEFORE_TITLE_PT)
    r = cap.add_run(f"{R.FIGURE_LABEL} {number}{R.FIGURE_CAPTION_DELIM}{title.strip()}")
    r.font.name = font
    r.font.size = Pt(R.FONT_SIZE_SMALL_PT)
    _add_bookmark(cap, f"{R.BM_FIGURE}{number}", counters)  # alvo da Lista de Ilustrações

    # Imagem centralizada (degrada para um aviso se o download/decodificação falhar)
    img = doc.add_paragraph()
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.paragraph_format.first_line_indent = Cm(0)
    run = img.add_run()
    try:
        stream = _load_image(url)
        shape = run.add_picture(stream)
        max_w = Cm(R.IMAGE_MAX_WIDTH_CM)
        if shape.width and shape.width > max_w:
            ratio = max_w / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)
    except Exception:
        run.text = "[imagem indisponível]"
        run.font.name = font
        run.font.size = Pt(R.FONT_SIZE_SMALL_PT)

    # Fonte na PARTE INFERIOR: "Fonte: ...", 10pt, centralizada
    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    src.paragraph_format.first_line_indent = Cm(0)
    src.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    src.paragraph_format.space_after = Pt(R.SPACE_AFTER_TITLE_PT)
    rs = src.add_run(_figure_source(fonte_raw, counters.get("year")))
    rs.font.name = font
    rs.font.size = Pt(R.FONT_SIZE_SMALL_PT)


def _figure_source(fonte_raw: str, year: str = None) -> str:
    return _source_line(fonte_raw, year, R.SOURCE_OWN_WORK)


def _source_line(fonte_raw: str, year: str, own_phrase: str) -> str:
    f = (fonte_raw or "").strip()
    if not f or f.lower() in _OWN_WORK_TOKENS:
        suffix = f" ({year})" if year else ""
        return f"Fonte: {own_phrase}{suffix}."
    if not f.lower().startswith("fonte"):
        f = "Fonte: " + f
    if not f.endswith("."):
        f += "."
    return f


def _load_image(url: str) -> io.BytesIO:
    url = url.strip()
    if url.startswith("data:"):
        _, _, b64 = url.partition(",")
        return io.BytesIO(base64.b64decode(b64))
    if url.startswith("http://") or url.startswith("https://"):
        req = urllib.request.Request(url, headers={"User-Agent": "ZaeonABNT/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:  # noqa: S310 (esquema validado acima)
            data = resp.read(R.IMAGE_MAX_BYTES + 1)
        if len(data) > R.IMAGE_MAX_BYTES:
            raise ValueError("imagem excede o tamanho máximo")
        return io.BytesIO(data)
    raise ValueError("esquema de URL de imagem não suportado")


# ── Tabelas (IBGE) e Quadros (NBR 14724:2024) ──────────────────
# Marcador no texto:
#   %%tabela Título da tabela
#   %%fonte dados da pesquisa (2026)        (opcional; "própria" → autor)
#   | Cabeçalho A | Cabeçalho B |
#   | --- | --- |                            (linha separadora opcional, ignorada)
#   | dado | dado |
#   (use %%quadro para um quadro fechado, dado qualitativo)
_TABLE_DIRECTIVE_RE = re.compile(r'^%%(tabela|quadro)\b', re.I)


def _is_table_block(text: str) -> bool:
    return bool(_TABLE_DIRECTIVE_RE.match(text.strip()))


def _parse_table_block(text: str):
    lines = text.strip().split("\n")
    m = re.match(r'^%%(tabela|quadro)\s*(.*)$', lines[0].strip(), re.I)
    kind = m.group(1).lower()
    title = m.group(2).strip()
    rest = lines[1:]

    fonte_raw = None
    if rest and re.match(r'^%%fonte\b', rest[0].strip(), re.I):
        fonte_raw = re.sub(r'^%%fonte\s*', '', rest[0].strip(), flags=re.I).strip()
        rest = rest[1:]

    rows = []
    for line in rest:
        s = line.strip()
        if '|' not in s:
            continue
        cells = [c.strip() for c in s.strip().strip('|').split('|')]
        non_empty = [c for c in cells if c]
        # ignora linha separadora markdown (ex.: | --- | :--: |)
        if non_empty and all(re.fullmatch(r':?-{2,}:?', c) for c in non_empty):
            continue
        rows.append(cells)
    return kind, title, fonte_raw, rows


def _add_table_block(doc: Document, text: str, font: str, counters: dict):
    kind, title, fonte_raw, rows = _parse_table_block(text)
    if not rows:
        return

    label = R.TABLE_LABEL if kind == "tabela" else R.QUADRO_LABEL
    counters[kind] = counters.get(kind, 0) + 1
    number = counters[kind]

    # Legenda no topo (acompanha a margem esquerda da tabela), 10pt
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cap.paragraph_format.space_before = Pt(R.SPACE_BEFORE_TITLE_PT)
    rc = cap.add_run(f"{label} {number}{R.FIGURE_CAPTION_DELIM}{title}")
    rc.font.name = font
    rc.font.size = Pt(R.FONT_SIZE_SMALL_PT)
    bm_prefix = R.BM_TABLE if kind == "tabela" else R.BM_QUADRO
    _add_bookmark(cap, f"{bm_prefix}{number}", counters)  # alvo da Lista de Tabelas/Ilustrações

    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if kind == "quadro":
        table.style = "Table Grid"  # caixa fechada (todas as bordas)

    for i, r in enumerate(rows):
        for j in range(ncols):
            value = r[j] if j < len(r) else ""
            _fill_cell(table.cell(i, j), value, font, bold=(i == 0))

    if kind == "tabela":
        _apply_ibge_borders(table)  # só traços horizontais (topo/cabeçalho/base)

    # Fonte no rodapé, 10pt
    own = R.SOURCE_OWN_WORK_F if kind == "tabela" else R.SOURCE_OWN_WORK
    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.LEFT
    src.paragraph_format.first_line_indent = Cm(0)
    src.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    src.paragraph_format.space_after = Pt(R.SPACE_AFTER_TITLE_PT)
    rs = src.add_run(_source_line(fonte_raw, counters.get("year"), own))
    rs.font.name = font
    rs.font.size = Pt(R.FONT_SIZE_SMALL_PT)


def _fill_cell(cell, value: str, font: str, bold: bool = False):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(value)
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_TABLE_PT)
    run.bold = bold


def _apply_ibge_borders(table):
    """Tabela IBGE: somente três traços horizontais (topo, sob o cabeçalho,
    base); nenhuma borda vertical/lateral."""
    for cell in table.rows[0].cells:
        _set_cell_borders(cell, top=True, bottom=True)   # topo + sob cabeçalho
    for cell in table.rows[-1].cells:
        _set_cell_borders(cell, bottom=True)             # base


def _set_cell_borders(cell, top: bool = False, bottom: bool = False):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge, on in (('top', top), ('bottom', bottom)):
        if not on:
            continue
        existing = tcBorders.find(qn(f'w:{edge}'))
        if existing is not None:
            tcBorders.remove(existing)
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        tcBorders.append(el)


# ── Referências ────────────────────────────────────────────────

def _add_references(doc: Document, references: list, font: str, counters: dict = None):
    tp = _section_title(doc, R.REFERENCES_TITLE, font, outline_level=0)
    if counters is not None:
        _add_bookmark(tp, R.BM_REFERENCES, counters)

    access_date = _today_abnt()
    ordered = sorted(references, key=_ref_sort_key)  # ordenação alfabética (NBR 6023)
    for ref in ordered:
        if getattr(ref, "title", None) or getattr(ref, "authors", None):
            _emit_structured_reference(doc, ref, font, access_date)
        else:
            _emit_text_reference(doc, ref.text, font)


def _new_reference_paragraph(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # referências à esquerda (NBR 6023)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE  # espaço simples interno
    p.paragraph_format.space_after = R.SPACE_BETWEEN_REFS          # espaço entre entradas
    return p


def _emit_text_reference(doc: Document, text: str, font: str):
    p = _new_reference_paragraph(doc)
    run = p.add_run(text or "")
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_BODY_PT)


def _emit_structured_reference(doc: Document, ref, font: str, access_date: str):
    """Formata uma referência conforme NBR 6023:2018 a partir de campos
    estruturados, com o DESTAQUE (negrito) no elemento correto: título da obra
    em monografias; título do periódico em artigos."""
    p = _new_reference_paragraph(doc)

    def run(text: str, bold: bool = False):
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(R.FONT_SIZE_BODY_PT)
        r.bold = bold

    authors = _abnt_authors(getattr(ref, "authors", []) or [])
    year = str(ref.year) if getattr(ref, "year", None) else "[s. d.]"
    title = (ref.title or "").strip().rstrip(".")
    venue = (ref.venue or "").strip().rstrip(".") if getattr(ref, "venue", None) else ""

    if authors:
        run(authors + (" " if authors.endswith(".") else ". "))  # evita "et al.."

    if venue:  # artigo de periódico: periódico em negrito, título do artigo normal
        if title:
            run(title + ". ")
        run(venue, bold=True)
        run(f", {year}.")
    else:      # monografia (livro/web/relatório): título em negrito
        if title:
            run(title, bold=True)
            run(". ")
        run(f"{year}.")

    if getattr(ref, "doi", None):
        run(f" DOI: {ref.doi.strip().rstrip('.')}.")
    elif getattr(ref, "url", None):
        ad = ref.access_date or access_date
        run(f" Disponível em: {ref.url.strip().rstrip('.')}. Acesso em: {ad}.")


# ── Formatação ABNT de autores (NBR 6023) ──────────────────────

def _abnt_author(name: str) -> str:
    """'João Silva' → 'SILVA, João'; 'Doudna J' → 'DOUDNA, J'."""
    tokens = name.strip().split()
    if not tokens:
        return ""
    if len(tokens) == 1:
        return tokens[0].upper()
    last = tokens[-1]
    stripped = last.replace(".", "")
    is_initials = (
        re.fullmatch(r"[A-Z]{1,3}", stripped) is not None
        and re.fullmatch(r"[A-Z.]{1,4}", last) is not None
        and len(tokens[0]) > 1
    )
    if is_initials:  # estilo "Sobrenome Iniciais" (PubMed)
        return f"{tokens[0].upper()}, {' '.join(tokens[1:])}"
    return f"{last.upper()}, {' '.join(tokens[:-1])}"


def _abnt_authors(authors: list) -> str:
    cleaned = [a.strip() for a in authors if a and a.strip()]
    if not cleaned:
        return ""
    if len(cleaned) > 3:  # 4+ autores → primeiro + et al.
        return f"{_abnt_author(cleaned[0])} et al."
    return "; ".join(_abnt_author(a) for a in cleaned)


def _ref_sort_key(ref) -> str:
    authors = getattr(ref, "authors", None)
    if authors:
        return _abnt_authors(authors).lower()
    if getattr(ref, "title", None):
        return ref.title.lower()
    return (getattr(ref, "text", "") or "").lower()


def _today_abnt() -> str:
    import datetime
    months = ["jan.", "fev.", "mar.", "abr.", "maio", "jun.",
              "jul.", "ago.", "set.", "out.", "nov.", "dez."]
    d = datetime.date.today()
    return f"{d.day} {months[d.month - 1]} {d.year}"


# ── Títulos ────────────────────────────────────────────────────

def _section_title(doc: Document, title: str, font: str, outline_level: int = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(R.SPACE_BEFORE_TITLE_PT)
    p.paragraph_format.space_after  = Pt(R.SPACE_AFTER_TITLE_PT)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_BODY_PT)
    if outline_level is not None:  # entra no sumário (somente seções do texto)
        _set_outline_level(p, outline_level)
    return p


def _subsection_title(doc: Document, title: str, font: str, outline_level: int = 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(R.SPACE_BEFORE_TITLE_PT)
    p.paragraph_format.space_after  = Pt(R.SPACE_AFTER_TITLE_PT)
    run = p.add_run(title.upper())
    run.bold = False   # subtítulos: sem negrito (NBR 14724)
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_BODY_PT)
    if outline_level is not None:
        _set_outline_level(p, outline_level)
    return p


def _blank_line(doc: Document, font: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    run = p.add_run("")
    run.font.name = font
    run.font.size = Pt(R.FONT_SIZE_BODY_PT)


# ── Numeração de páginas ───────────────────────────────────────

def _add_page_numbers(doc: Document):
    """Numeração no canto superior direito, visível SOMENTE na parte textual
    (última seção). Capa e pré-textuais ficam sem número (NBR 14724)."""
    last = len(doc.sections) - 1
    for i, section in enumerate(doc.sections):
        header = section.header
        header.is_linked_to_previous = False
        for para in header.paragraphs:
            para.clear()
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)

        if i != last:
            continue  # capa e pré-textuais: header vazio (sem número)

        run = p.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'PAGE'
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
