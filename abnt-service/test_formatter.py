"""
Testes de conformidade ABNT NBR 14724:2011 do motor de formatação.

Constroem um .docx real via build_abnt_document e inspecionam o resultado com
python-docx, garantindo que as regras da norma estão aplicadas (margens,
fonte, entrelinha, justificado, citações curtas×longas, sumário, abstract,
referências).
"""
import io

import pytest
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

import abnt_rules as R
from formatter import build_abnt_document
from models import (
    DocumentPayload,
    CoverData,
    SectionData,
    SubsectionData,
    ReferenceData,
    AppendixData,
)


# ── Helpers ────────────────────────────────────────────────────

BODY_MARKER = "CORPO_DE_TESTE_PARAGRAFO_NORMAL que precisa ser justificado."
SHORT_QUOTE = '"Uma citação curta cabe no corpo."'
# Citação longa: > 3 linhas (>~270 chars), iniciada por aspas.
LONG_QUOTE = (
    '"' + ("Esta é uma citação direta longa que ultrapassa tres linhas e "
           "portanto deve ser formatada em bloco com recuo de quatro centimetros, "
           "fonte menor de dez pontos e espacamento simples conforme a norma "
           "ABNT NBR 10520 exige para citacoes longas em trabalhos academicos. ") * 2 + '"'
)
# Prosa longa SEM aspas: NÃO deve virar citação longa (deve ficar justificada no corpo).
LONG_PROSE = (
    "Este parágrafo de prosa comum também ultrapassa três linhas de texto, "
    "mas não é uma citação e portanto deve permanecer como corpo justificado "
    "normal, com recuo de primeira linha, e jamais ser transformado em bloco "
    "de citação recuado e em fonte reduzida, o que seria um erro grave."
)


def _build(**overrides) -> Document:
    payload_kwargs = dict(
        title="Trabalho de Teste",
        work_type="Artigo",
        cover=CoverData(
            author="Fulano de Tal",
            institution="Universidade Federal de Teste",
            course="Engenharia",
            advisor="Prof. Ciclano",
            city="São Paulo",
            year="2026",
        ),
        resumo="Este é o resumo do trabalho em português para fins de teste.",
        palavras_chave=["ABNT", "Formatação", "Teste"],
        abstract="This is the English abstract of the work for testing purposes.",
        keywords=["ABNT", "Formatting", "Test"],
        sections=[
            SectionData(
                type="introducao",
                title="Introdução",
                content=BODY_MARKER + "\n\n" + SHORT_QUOTE + "\n\n" + LONG_QUOTE
                + "\n\n" + LONG_PROSE,
                subsections=[
                    SubsectionData(title="Contexto", content="Texto da subseção.")
                ],
            ),
            SectionData(
                type="conclusao", title="Conclusão", content="Considerações finais.",
            ),
        ],
        references=[
            ReferenceData(text="SOBRENOME, Nome. Título da obra. Cidade: Editora, 2020."),
            ReferenceData(text="OUTRO, Autor. Segunda obra. Cidade: Editora, 2021."),
        ],
        format="docx",
        font_choice="Times New Roman",
    )
    payload_kwargs.update(overrides)
    data = build_abnt_document(DocumentPayload(**payload_kwargs))
    return Document(io.BytesIO(data))


def _find_para(doc, needle):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    return None


def _body_xml(doc) -> str:
    return doc.element.body.xml


def cm_eq(length, cm_value) -> bool:
    """Word grava margens/recuos em twips, então Cm(3) volta como ~3.0004cm.
    Compara em cm com 2 casas para absorver o arredondamento."""
    if length is None:
        return round(cm_value, 2) == 0.0
    return round(length.cm, 2) == round(cm_value, 2)


# ── Página / margens (regressão — já implementado) ─────────────

def test_page_margins_abnt():
    doc = _build()
    sec = doc.sections[0]
    assert cm_eq(sec.left_margin, 3.0)
    assert cm_eq(sec.top_margin, 3.0)
    assert cm_eq(sec.right_margin, 2.0)
    assert cm_eq(sec.bottom_margin, 2.0)


def test_base_font_and_spacing():
    doc = _build()
    normal = doc.styles["Normal"]
    assert normal.font.size == Pt(12)
    assert normal.font.name == "Times New Roman"
    assert normal.paragraph_format.line_spacing == 1.5


def test_font_choice_arial_propagates():
    doc = _build(font_choice="Arial")
    assert doc.styles["Normal"].font.name == "Arial"


# ── Justificado (NOVO) ─────────────────────────────────────────

def test_body_paragraph_is_justified():
    doc = _build()
    p = _find_para(doc, BODY_MARKER)
    assert p is not None, "parágrafo de corpo não encontrado"
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_resumo_paragraph_is_justified():
    doc = _build()
    p = _find_para(doc, "resumo do trabalho em português")
    assert p is not None
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_references_aligned_left_explicitly():
    doc = _build()
    p = _find_para(doc, "SOBRENOME, Nome.")
    assert p is not None
    assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT


# ── Citações curta × longa por nº de linhas (NOVO) ─────────────

def test_long_quote_block_formatting():
    doc = _build()
    # citação longa: recuo 4cm, fonte 10pt, sem aspas envolventes
    p = None
    for cand in doc.paragraphs:
        if "citação direta longa que ultrapassa" in cand.text:
            p = cand
            break
    assert p is not None, "citação longa não encontrada"
    assert cm_eq(p.paragraph_format.left_indent, R.LONG_QUOTE_INDENT_CM)
    assert cm_eq(p.paragraph_format.first_line_indent, 0)
    assert p.runs and p.runs[0].font.size == Pt(R.FONT_SIZE_QUOTE_PT)
    # aspas envolventes removidas no bloco
    assert not p.text.strip().startswith('"')


def test_short_quote_stays_inline_body():
    doc = _build()
    p = _find_para(doc, "Uma citação curta cabe no corpo")
    assert p is not None
    # curta: permanece corpo (12pt, sem recuo de 4cm)
    assert cm_eq(p.paragraph_format.left_indent, 0)
    assert p.runs[0].font.size == Pt(R.FONT_SIZE_BODY_PT)


def test_long_prose_not_treated_as_quote():
    doc = _build()
    p = _find_para(doc, "parágrafo de prosa comum também ultrapassa")
    assert p is not None
    # prosa longa sem aspas: corpo justificado, recuo de 1ª linha, 12pt
    assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert cm_eq(p.paragraph_format.first_line_indent, R.PARAGRAPH_INDENT_CM)
    assert p.runs[0].font.size == Pt(R.FONT_SIZE_BODY_PT)


# ── Sumário real (NOVO) ────────────────────────────────────────

def test_sumario_entries_use_pageref():
    doc = _build()
    xml = _body_xml(doc)
    # números de página via PAGEREF (resolvidos no layout pelo Word/LibreOffice)
    assert "PAGEREF" in xml, "sumário deve usar PAGEREF para números de página"
    assert "Atualize o sum" not in xml, "não deve vazar texto-instrução"
    # bookmarks-alvo presentes nos títulos
    assert R.BM_SECTION + "1" in xml and R.BM_REFERENCES in xml


def test_section_titles_have_bookmark():
    doc = _build()
    p = next((p for p in doc.paragraphs if p.text.strip() == "1 INTRODUÇÃO"), None)
    assert p is not None
    assert "bookmarkStart" in p._p.xml, "título de seção precisa de bookmark (alvo do PAGEREF)"


def test_sumario_cached_entries_present():
    """O resultado em cache do campo TOC já lista as seções, para que o PDF
    headless exiba a estrutura do sumário (e não uma página em branco)."""
    doc = _build()
    texts = [p.text for p in doc.paragraphs]
    # "1 INTRODUÇÃO" aparece como título do corpo E como entrada do sumário
    intro_count = sum(1 for t in texts if "1 INTRODUÇÃO" in t)
    assert intro_count >= 2, "entradas em cache do sumário ausentes"
    # subseção e referências também entram no sumário (título do corpo + entrada)
    assert sum(1 for t in texts if "1.1 CONTEXTO" in t) >= 2
    assert sum(1 for t in texts if "REFERÊNCIAS" in t) >= 2


def test_section_titles_have_outline_level():
    doc = _build()
    # "1 INTRODUÇÃO" aparece como entrada do sumário (sem nível) e como título
    # do corpo (com nível); ao menos o título do corpo deve ter o outlineLvl.
    matches = [p for p in doc.paragraphs if "1 INTRODUÇÃO" in p.text]
    assert matches, "título de seção não encontrado"
    assert any("w:outlineLvl" in p._p.xml for p in matches), \
        "título de seção sem nível de tópico (não entra no sumário)"


# ── Abstract (NOVO) ────────────────────────────────────────────

def test_abstract_section_present():
    doc = _build()
    p = _find_para(doc, "ABSTRACT")
    assert p is not None, "seção ABSTRACT ausente"


def test_abstract_keywords_line_present():
    doc = _build()
    p = _find_para(doc, "Keywords")
    assert p is not None, "linha de Keywords (abstract) ausente"


def test_no_abstract_when_absent():
    doc = _build(abstract=None, keywords=[])
    assert _find_para(doc, "ABSTRACT") is None


# ── Figuras / Ilustrações (NBR 14724:2024 + IBGE) — NOVO ───────

# PNG 4x4 válido em data URI (evita rede nos testes)
FIG_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAQAAAAECAIAAAAmkwkp"
           "AAAAEElEQVR4nGMwTpsJRwzEcQDrIxMhwmPP/AAAAABJRU5ErkJggg==")


def _fig_doc(content):
    return _build(sections=[SectionData(type="desenvolvimento", title="Desenvolvimento",
                                        content=content)])


def test_figure_caption_top_image_and_source_bottom():
    doc = _fig_doc(f"Veja a figura a seguir.\n\n"
                   f"![Mapa do Brasil|Fonte: IBGE (2020)]({FIG_PNG})\n\nTexto após.")
    # legenda do CORPO (centralizada, 10pt) — não a entrada da Lista de Ilustrações
    cap = next((p for p in doc.paragraphs
                if "Figura 1 – Mapa do Brasil" in p.text
                and p.alignment == WD_ALIGN_PARAGRAPH.CENTER), None)
    assert cap is not None, "legenda da figura ausente"
    assert cap.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert cap.runs[0].font.size == Pt(R.FONT_SIZE_SMALL_PT)  # 10pt
    fonte = _find_para(doc, "Fonte: IBGE (2020)")
    assert fonte is not None, "linha de fonte ausente"
    assert fonte.runs[0].font.size == Pt(R.FONT_SIZE_SMALL_PT)
    assert len(doc.inline_shapes) >= 1, "imagem não inserida"


def test_figure_caption_uses_en_dash_not_hyphen():
    doc = _fig_doc(f"![Título qualquer|própria]({FIG_PNG})")
    cap = _find_para(doc, "Figura 1")
    assert cap is not None
    assert "–" in cap.text, "deve usar travessão (en-dash), não hífen"


def test_figure_own_work_source_default():
    doc = _fig_doc(f"![Gráfico de barras|própria]({FIG_PNG})")
    assert _find_para(doc, "elaborado pelo próprio autor") is not None


def test_figure_numbering_is_sequential():
    doc = _fig_doc(f"![Primeira|própria]({FIG_PNG})\n\n![Segunda|própria]({FIG_PNG})")
    assert _find_para(doc, "Figura 1 – Primeira") is not None
    assert _find_para(doc, "Figura 2 – Segunda") is not None


def test_figure_marker_not_left_as_raw_text():
    doc = _fig_doc(f"![Algo|própria]({FIG_PNG})")
    assert not any("![" in p.text for p in doc.paragraphs), "marcador markdown vazou como texto"


# ── Tabelas (IBGE) e Quadros — NOVO ────────────────────────────

TABELA = ("%%tabela Distribuição de alunos por turma\n"
          "%%fonte dados da pesquisa (2026)\n"
          "| Turma | Alunos |\n"
          "| --- | --- |\n"
          "| A | 30 |\n"
          "| B | 25 |")

QUADRO = ("%%quadro Comparação entre abordagens\n"
          "%%fonte própria\n"
          "| Critério | Método X | Método Y |\n"
          "| Custo | Alto | Baixo |")


def _tab_doc(content):
    return _build(sections=[SectionData(type="desenvolvimento", title="Resultados",
                                        content=content)])


def _cell_edges(cell):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is None:
        return set()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        return set()
    return {c.tag.split('}')[-1] for c in b if c.get(qn('w:val')) not in (None, 'nil', 'none')}


def test_tabela_caption_top_source_bottom():
    doc = _tab_doc(TABELA)
    # legenda do CORPO (10pt) — não a entrada da Lista de Tabelas (12pt)
    cap = next((p for p in doc.paragraphs
                if "Tabela 1 – Distribuição de alunos por turma" in p.text
                and p.runs and p.runs[0].font.size == Pt(R.FONT_SIZE_SMALL_PT)), None)
    assert cap is not None, "legenda da tabela ausente"
    assert cap.runs[0].font.size == Pt(R.FONT_SIZE_SMALL_PT)
    assert "–" in cap.text  # travessão
    assert _find_para(doc, "Fonte: dados da pesquisa (2026)") is not None
    assert len(doc.tables) == 1
    # célula contém o dado
    assert any("30" in c.text for row in doc.tables[0].rows for c in row.cells)


def test_tabela_ibge_no_vertical_borders_three_rules():
    doc = _tab_doc(TABELA)
    t = doc.tables[0]
    # nenhuma célula pode ter borda esquerda/direita (tabela aberta lateralmente)
    for row in t.rows:
        for c in row.cells:
            edges = _cell_edges(c)
            assert "left" not in edges and "right" not in edges, "tabela não pode ter bordas verticais"
    header = t.rows[0].cells
    assert all("top" in _cell_edges(c) for c in header), "falta traço no topo"
    assert all("bottom" in _cell_edges(c) for c in header), "falta traço sob o cabeçalho"
    assert all("bottom" in _cell_edges(c) for c in t.rows[-1].cells), "falta traço na base"


def test_tabela_centered():
    doc = _tab_doc(TABELA)
    assert doc.tables[0].alignment == WD_TABLE_ALIGNMENT.CENTER


def test_quadro_closed_borders_and_label():
    doc = _tab_doc(QUADRO)
    assert _find_para(doc, "Quadro 1 – Comparação entre abordagens") is not None
    assert _find_para(doc, "elaborado pelo próprio autor") is not None
    t = doc.tables[0]
    assert t.style is not None and t.style.name == "Table Grid", "quadro deve ser caixa fechada"


def test_tabela_quadro_independent_numbering():
    doc = _tab_doc(TABELA + "\n\n" + QUADRO + "\n\n" + TABELA)
    assert _find_para(doc, "Tabela 1 – Distribuição de alunos por turma") is not None
    assert _find_para(doc, "Quadro 1 – Comparação entre abordagens") is not None
    assert _find_para(doc, "Tabela 2 – Distribuição de alunos por turma") is not None


def test_list_of_illustrations_with_figures_and_quadros():
    content = (f"![Mapa do Brasil|própria]({FIG_PNG})\n\n"
               f"%%quadro Comparação de métodos\n| C | X | Y |\n| Custo | A | B |")
    doc = _build(sections=[SectionData(type="d", title="Dev", content=content)])
    assert _find_para(doc, "LISTA DE ILUSTRAÇÕES") is not None
    texts = [p.text for p in doc.paragraphs]
    assert any("Figura 1 – Mapa do Brasil" in t for t in texts)
    assert any("Quadro 1 – Comparação de métodos" in t for t in texts)
    assert "PAGEREF" in _body_xml(doc), "lista deve usar PAGEREF"
    caps = [p for p in doc.paragraphs if "Figura 1 – Mapa do Brasil" in p.text]
    assert any("bookmarkStart" in p._p.xml for p in caps), "legenda da figura precisa de bookmark"


def test_list_of_tables_present():
    content = "%%tabela Dados coletados\n| Ano | Valor |\n| 2020 | 10 |"
    doc = _build(sections=[SectionData(type="d", title="Dev", content=content)])
    assert _find_para(doc, "LISTA DE TABELAS") is not None
    assert any("Tabela 1 – Dados coletados" in p.text for p in doc.paragraphs)


def test_no_lists_when_no_floats():
    doc = _build(sections=[SectionData(type="i", title="Intro", content="Apenas texto comum.")])
    assert _find_para(doc, "LISTA DE ILUSTRAÇÕES") is None
    assert _find_para(doc, "LISTA DE TABELAS") is None


def test_table_marker_not_left_as_raw_text():
    doc = _tab_doc(TABELA)
    assert not any("%%tabela" in p.text or "%%fonte" in p.text for p in doc.paragraphs)
    assert not any(p.text.strip().startswith("|") for p in doc.paragraphs), "linha markdown vazou"


# ── Referências determinísticas (NBR 6023:2018) — NOVO ─────────

def _bold_runs(p):
    return [r.text for r in p.runs if r.bold]


def test_reference_monograph_title_bold():
    # sem venue → monografia → TÍTULO em negrito
    ref = ReferenceData(text="fb", authors=["Maria Silva Santos"],
                        title="A história do Brasil", year=2020)
    doc = _build(references=[ref])
    p = _find_para(doc, "SANTOS, Maria Silva")
    assert p is not None, "autor não formatado (SOBRENOME, Prenome)"
    assert any("A história do Brasil" in t for t in _bold_runs(p)), "título deveria estar em negrito"
    assert "2020" in p.text


def test_reference_article_venue_bold_not_title():
    # com venue → artigo de periódico → PERIÓDICO em negrito, título do artigo não
    ref = ReferenceData(text="fb", authors=["João Pereira", "Ana Lima"],
                        title="Um estudo sobre IA", year=2021,
                        venue="Revista Brasileira de Computação", doi="10.1234/abc")
    doc = _build(references=[ref])
    p = _find_para(doc, "PEREIRA, João")
    assert p is not None
    bold = _bold_runs(p)
    assert any("Revista Brasileira de Computação" in t for t in bold), "periódico deveria estar em negrito"
    assert not any("Um estudo sobre IA" in t for t in bold), "título do artigo NÃO leva negrito"
    assert "DOI: 10.1234/abc" in p.text


def test_reference_online_disponivel_acesso():
    ref = ReferenceData(text="fb", authors=["Carlos Souza"],
                        title="Página de exemplo", year=2019, url="https://exemplo.com/x")
    doc = _build(references=[ref])
    p = _find_para(doc, "SOUZA, Carlos")
    assert "Disponível em: https://exemplo.com/x" in p.text
    assert "Acesso em:" in p.text


def test_reference_four_authors_et_al():
    ref = ReferenceData(text="fb",
                        authors=["Ana Silva", "Bruno Costa", "Carla Dias", "Diego Reis"],
                        title="Obra coletiva", year=2022)
    doc = _build(references=[ref])
    p = _find_para(doc, "et al.")
    assert p is not None
    assert "et al.." not in p.text, "ponto duplicado após et al."


def test_reference_fallback_to_text_when_unstructured():
    ref = ReferenceData(text="TEXTO PRE FORMATADO COMPLETO ABNT.")
    doc = _build(references=[ref])
    assert _find_para(doc, "TEXTO PRE FORMATADO COMPLETO ABNT.") is not None


def test_appendices_and_annexes_rendered():
    doc = _build(
        apendices=[AppendixData(title="Questionário aplicado", content="Conteúdo do questionário.")],
        anexos=[AppendixData(title="Parecer do comitê", content="Texto do parecer.")],
    )
    assert _find_para(doc, "APÊNDICE A – QUESTIONÁRIO APLICADO") is not None
    assert _find_para(doc, "Conteúdo do questionário.") is not None
    assert _find_para(doc, "ANEXO A – PARECER DO COMITÊ") is not None
    assert _find_para(doc, "Texto do parecer.") is not None


def test_appendices_appear_in_sumario():
    doc = _build(apendices=[AppendixData(title="Questionário", content="x")])
    texts = [p.text for p in doc.paragraphs]
    # entrada no sumário + título no corpo
    assert sum(1 for t in texts if "APÊNDICE A" in t) >= 2


def test_alineas_lettered_with_punctuation():
    content = ("Os objetivos específicos são:\n\n"
               "- analisar os dados coletados\n"
               "- propor um modelo conceitual\n"
               "- validar os resultados obtidos")
    doc = _build(sections=[SectionData(type="d", title="Dev", content=content)])
    texts = [p.text for p in doc.paragraphs]
    assert "a) analisar os dados coletados;" in texts
    assert "b) propor um modelo conceitual;" in texts
    assert "c) validar os resultados obtidos." in texts  # última com ponto final


def test_alineas_not_triggered_for_inline_dash():
    content = "Este texto tem um - hífen no meio e não é uma lista de alíneas."
    doc = _build(sections=[SectionData(type="d", title="Dev", content=content)])
    assert not any(p.text.strip().startswith("a)") for p in doc.paragraphs)


def _footnotes_blob(doc):
    for p in doc.part.package.iter_parts():
        if "footnotes.xml" in str(p.partname):
            return p.blob.decode("utf-8")
    return None


def test_footnote_creates_part_and_reference():
    content = "Afirmação importante^[Esta é uma nota explicativa de rodapé.] seguida de texto."
    doc = _build(sections=[SectionData(type="d", title="Dev", content=content)])
    blob = _footnotes_blob(doc)
    assert blob is not None, "parte footnotes.xml ausente"
    assert "footnoteReference" in doc.element.body.xml, "referência de nota ausente no corpo"
    assert "Esta é uma nota explicativa de rodapé." in blob


def test_footnote_text_is_10pt():
    doc = _build(sections=[SectionData(type="d", title="Dev", content="X^[nota curta]")])
    blob = _footnotes_blob(doc)
    assert f'w:sz w:val="{R.FONT_SIZE_NOTE_PT * 2}"' in blob  # 10pt = 20 meios-pontos


def test_footnote_marker_not_left_as_raw_text():
    doc = _build(sections=[SectionData(type="d", title="Dev", content="Texto^[nota] e mais.")])
    assert not any("^[" in p.text for p in doc.paragraphs), "marcador de nota vazou como texto"


def test_no_footnotes_part_when_absent():
    doc = _build(sections=[SectionData(type="d", title="Dev", content="Sem notas aqui.")])
    assert _footnotes_blob(doc) is None


def _header_xml(section):
    paras = section.header.paragraphs
    return paras[0]._p.xml if paras else ""


def test_pagination_three_sections():
    # capa | pré-textual | textual+pós-textual
    doc = _build()
    assert len(doc.sections) == 3


def test_pretextual_section_restarts_numbering_at_1():
    # folha de rosto = página 1 (capa não é contada)
    doc = _build()
    sectPr = doc.sections[1]._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    assert pg is not None and pg.get(qn('w:start')) == '1', "pré-textual deve reiniciar em 1"


def test_textual_section_continues_numbering_not_restart():
    # a parte textual NÃO pode reiniciar em 1 (intro deve mostrar a contagem acumulada)
    doc = _build()
    pg = doc.sections[2]._sectPr.find(qn('w:pgNumType'))
    start = pg.get(qn('w:start')) if pg is not None else None
    assert start is None, "seção textual não pode reiniciar a numeração"


def test_page_number_only_visible_in_textual_section():
    doc = _build()
    assert 'PAGE' not in _header_xml(doc.sections[0]), "capa não pode ter número"
    assert 'PAGE' not in _header_xml(doc.sections[1]), "pré-textual não exibe número"
    assert 'PAGE' in _header_xml(doc.sections[2]), "texto deve exibir número"


def test_textual_page_number_top_right():
    doc = _build()
    p = doc.sections[2].header.paragraphs[0]
    assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_all_sections_keep_abnt_margins():
    doc = _build()
    for s in doc.sections:
        assert cm_eq(s.left_margin, 3.0) and cm_eq(s.top_margin, 3.0)
        assert cm_eq(s.right_margin, 2.0) and cm_eq(s.bottom_margin, 2.0)


def test_reference_alphabetical_order():
    refs = [
        ReferenceData(text="fb", authors=["Zeca Ulisses"], title="Zeta", year=2020),
        ReferenceData(text="fb", authors=["Ana Alves"], title="Alfa", year=2020),
    ]
    doc = _build(references=refs)
    texts = [p.text for p in doc.paragraphs]
    i_alves = next(i for i, t in enumerate(texts) if "ALVES, Ana" in t)
    i_ulisses = next(i for i, t in enumerate(texts) if "ULISSES, Zeca" in t)
    assert i_alves < i_ulisses, "referências devem estar em ordem alfabética"
